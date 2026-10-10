# -*- coding: utf-8 -*-
"""
pptx_kit —— 中文商务 PPT（python-pptx）通用设计工具包
用法：from pptx_kit import *
      prs = new_deck(); s = slide(prs, dark=True); ...
设计原则：16:9 / 米纸白底 + 深色强调页 / 卡片+细描边 / 强调色方块点缀 / 页眉页脚页码系统
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

W, H = 13.333, 7.5
MG = 0.78
FONT = '微软雅黑'

# ---------- 品牌配色（改这里即可换肤） ----------
PALETTES = {
    # 摘柠檬：跨境电商线 / 跨境101
    'takelemon': dict(paper='FAF7EE', ink='191712', ink_tint='2A271F', accent='FFD60A',
                      accent2='FFB347', accent_text='C79A00', gray='6F6A5C', gray_l='9C9686',
                      card='FFFFFF', border='E6DFC9', soft='FFF6CE', paper_soft='E9E3D6'),
    # 靛蓝瓷：教育控股 / 稳重商务
    'indigo': dict(paper='F1F3F5', ink='0A1F3D', ink_tint='152A4A', accent='2F6FED',
                   accent2='6FA8FF', accent_text='1B5ACB', gray='5B6472', gray_l='98A1AF',
                   card='FFFFFF', border='DCE3EC', soft='E8F0FF', paper_soft='D8E0EA'),
    # 深空黑蓝 + 重启橙：AI 科技公司
    'reboot': dict(paper='F4F6F9', ink='0B1220', ink_tint='17203A', accent='FF6B2C',
                   accent2='FFA46B', accent_text='D9521A', gray='5C6577', gray_l='98A1B2',
                   card='FFFFFF', border='E2E7EF', soft='FFF0E8', paper_soft='DDE4EE'),
}


def use_palette(name='takelemon'):
    """把配色写入模块级常量，返回 dict"""
    global PAPER, INK, INK_TINT, LEMON, AMBER, GOLD, GRAY, GRAY_L, CARD, BORDER, LEMON_SOFT, PAPER_SOFT
    p = PALETTES[name]
    RGB = lambda k: RGBColor.from_string(p[k])
    PAPER, INK, INK_TINT = RGB('paper'), RGB('ink'), RGB('ink_tint')
    LEMON, AMBER, GOLD = RGB('accent'), RGB('accent2'), RGB('accent_text')
    GRAY, GRAY_L = RGB('gray'), RGB('gray_l')
    CARD, BORDER, LEMON_SOFT, PAPER_SOFT = RGB('card'), RGB('border'), RGB('soft'), RGB('paper_soft')
    return p


use_palette('takelemon')


def set_font(run, font=FONT):
    """同时设置西文与中文字体（只设 run.font.name 中文会回落到宋体）"""
    run.font.name = font
    rPr = run.font._rPr
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', font)
    return run


def new_deck(w=W, h=H):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(w), Inches(h)
    prs._deck_w, prs._deck_h = w, h
    return prs


def slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = rect(s, 0, 0, prs.slide_width / 914400, prs.slide_height / 914400, fill=INK if dark else PAPER)
    r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size, bold=False, color=INK, first=False, align=PP_ALIGN.LEFT,
         spacing=1.15, before=0, after=0, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment, p.line_spacing = align, spacing
    p.space_before, p.space_after = Pt(before), Pt(after)
    r = p.add_run(); r.text = text
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    set_font(r, font)
    return p


def rich(tf, segs, size, first=False, align=PP_ALIGN.LEFT, spacing=1.15, before=0, after=0, font=FONT):
    """segs = [(text, bold, color), ...] 一句话内混色"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment, p.line_spacing = align, spacing
    p.space_before, p.space_after = Pt(before), Pt(after)
    for text, bold, color in segs:
        r = p.add_run(); r.text = text
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        set_font(r, font)
    return p


def rule(s, x, y, w, color=BORDER, h=0.012):
    return rect(s, x, y, w, h, fill=color)


def chip(s, x, y, size=0.11, color=LEMON):
    return rect(s, x, y, size, size, fill=color)


def header(s, no, right_text, dark=False, brand='', total=0):
    c = PAPER_SOFT if dark else GRAY
    para(tbox(s, MG, 0.4, 7.0, 0.3), brand, 9, False, c, first=True)
    para(tbox(s, W - MG - 5.0, 0.4, 5.0, 0.3), right_text, 8, False, c, first=True, align=PP_ALIGN.RIGHT)
    rule(s, MG, 0.78, W - 2 * MG, INK_TINT if dark else BORDER, 0.01)


def footer(s, no, left_text, dark=False, total=0):
    c = GRAY_L if dark else GRAY
    rule(s, MG, 6.83, W - 2 * MG, INK_TINT if dark else BORDER, 0.01)
    para(tbox(s, MG, 6.94, 8.0, 0.3), left_text, 8, False, c, first=True)
    if total:
        para(tbox(s, W - MG - 3.0, 6.94, 3.0, 0.3), '%d / %d' % (no, total), 8, False, c,
             first=True, align=PP_ALIGN.RIGHT)


def kicker(s, x, y, text, dark=False):
    """强调色小方块 + 小标题，统一视觉锚点"""
    chip(s, x, y + 0.045, 0.11, 0.11)
    para(tbox(s, x + 0.24, y, 9.0, 0.3), text, 10.5, True, PAPER_SOFT if dark else GOLD, first=True)


def title(s, text, y=1.45, size=34, dark=False, x=MG, w=11.0):
    para(tbox(s, x, y, w, 0.8), text, size, True, PAPER if dark else INK, first=True)


def picture(s, path, x=None, y=None, w=1.25):
    if os.path.exists(path):
        x = W - MG - w if x is None else x
        y = H - 1.55 if y is None else y
        s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
